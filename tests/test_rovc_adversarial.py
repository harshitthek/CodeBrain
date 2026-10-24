"""Adversarial Verification Pytest Suite for DTU ROVC 2026.

Empirically challenges:
1. Pydantic schema validation & adversarial mutations
2. Coverage matrix (Tasks 1-3, Safety, Tiers)
3. Boundary thresholds (35.00kg vs 35.05kg, 1.0m width, 0.6m/1.8m depth, L<75cm)
4. Exception & fault scenarios
5. TestCaseReviewerAgent.evaluate_suite sensitivity & discrimination
6. Multi-format exporter security & integrity (Excel injection sanitization, XML, Mindmap)
"""

import copy
import json
import os
import re
import xml.etree.ElementTree as ET
from pathlib import Path
import pytest
import openpyxl
from pydantic import ValidationError

from apps.ai_agents.test_case_generator.schemas import GeneratedTestCaseItem
from apps.ai_agents.test_case_reviewer.reviewer import TestCaseReviewerAgent
from apps.core.exporters.excel_exporter import _sanitize_excel_text

BASE_DIR = Path(__file__).resolve().parent.parent


def _resolve_export_dir() -> Path:
    candidates = []
    if os.getenv("TUTORIAL_DIR"):
        candidates.append(Path(os.getenv("TUTORIAL_DIR")))
    candidates.append(BASE_DIR.parent / "Tutorial")
    candidates.append(BASE_DIR / "docs" / "tutorial_qa")
    candidates.append(Path("/Users/pernavjain/Documents/Tutorial"))
    for c in candidates:
        if (c / "rovc_testcases.json").exists():
            return c
    return BASE_DIR / "docs" / "tutorial_qa"


TUTORIAL_DIR = _resolve_export_dir()
JSON_PATH = TUTORIAL_DIR / "rovc_testcases.json"
XLSX_PATH = TUTORIAL_DIR / "rovc_testcases.xlsx"
XML_PATH = TUTORIAL_DIR / "junit_testcases.xml"
MINDMAP_PATH = TUTORIAL_DIR / "testcases_mindmap.md"


@pytest.fixture(scope="module")
def rovc_data():
    assert JSON_PATH.exists(), f"Missing {JSON_PATH}"
    with open(JSON_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


@pytest.fixture(scope="module")
def test_cases(rovc_data):
    return rovc_data.get("test_cases", [])


def test_adversarial_schema_validation_and_mutations(test_cases):
    """Validate all 40 cases against GeneratedTestCaseItem and challenge schema with adversarial mutations."""
    assert len(test_cases) == 40
    case_id_pattern = re.compile(r"^ROVC-TC-\d{3}$")
    seen_ids = set()

    for idx, tc in enumerate(test_cases, start=1):
        validated = GeneratedTestCaseItem.model_validate(tc)
        dump = validated.model_dump()
        assert dump["case_id"] == f"ROVC-TC-{idx:03d}"
        assert case_id_pattern.match(dump["case_id"])
        assert dump["case_id"] not in seen_ids
        seen_ids.add(dump["case_id"])

        assert len(dump["name"].strip()) >= 10
        assert len(dump["description"].strip()) >= 20
        assert len(dump["preconditions"].strip()) >= 15
        assert isinstance(dump["test_steps"], list) and len(dump["test_steps"]) >= 2
        assert isinstance(dump["expected_results"], list) and len(dump["expected_results"]) >= 1
        assert dump["case_type"] in ["Functional", "Boundary", "Exception", "Security", "Performance", "UI"]
        assert dump["priority"] in ["High", "Medium", "Low", "P0", "P1", "P2", "P3"]

    # Adversarial Mutation 1: Empty description rejected
    with pytest.raises(ValidationError):
        GeneratedTestCaseItem.model_validate({
            "case_id": "ROVC-TC-MUT1",
            "description": "",
            "test_steps": ["step 1", "step 2"],
            "expected_results": ["result 1"]
        })

    # Adversarial Mutation 2: Empty test steps rejected
    with pytest.raises(ValidationError):
        GeneratedTestCaseItem.model_validate({
            "case_id": "ROVC-TC-MUT2",
            "description": "Valid description scenario here",
            "test_steps": [],
            "expected_results": ["result 1"]
        })

    # Adversarial Mutation 3: Empty expected results rejected
    with pytest.raises(ValidationError):
        GeneratedTestCaseItem.model_validate({
            "case_id": "ROVC-TC-MUT3",
            "description": "Valid description scenario here",
            "test_steps": ["step 1", "step 2"],
            "expected_results": []
        })


def test_adversarial_coverage_matrix(test_cases):
    """Empirically assert presence of Tasks 1-3, physical/safety constraints, and participant categories."""
    all_text = " ".join([
        f"{tc.get('name', '')} {tc.get('description', '')} {' '.join(tc.get('test_steps', []))} {' '.join(tc.get('expected_results', []))}"
        for tc in test_cases
    ]).lower()

    # Tasks 1-3
    t1 = [tc for tc in test_cases if "Task 1" in tc.get("domain", "") or "gate" in (tc.get("name", "") + " " + tc.get("description", "")).lower()]
    assert len(t1) >= 6

    t2 = [tc for tc in test_cases if "Task 2" in tc.get("domain", "") or any(k in (tc.get("name", "") + " " + tc.get("description", "")).lower() for k in ["vision", "balloon", "touch"])]
    assert len(t2) >= 7

    t3 = [tc for tc in test_cases if "Task 3" in tc.get("domain", "") or "maze" in (tc.get("name", "") + " " + tc.get("description", "")).lower()]
    assert len(t3) >= 8

    # Safety & Physical
    assert any("weight" in (tc.get("name", "") + " " + tc.get("description", "")).lower() for tc in test_cases)
    assert any("kill switch" in (tc.get("name", "") + " " + tc.get("description", "")).lower() for tc in test_cases)
    assert any("fuse" in (tc.get("name", "") + " " + tc.get("description", "")).lower() for tc in test_cases)
    assert any("footprint" in (tc.get("name", "") + " " + tc.get("description", "")).lower() or "envelope" in (tc.get("name", "") + " " + tc.get("description", "")).lower() for tc in test_cases)
    assert any("pneumatic" in (tc.get("name", "") + " " + tc.get("description", "")).lower() for tc in test_cases)
    assert any("leak" in (tc.get("name", "") + " " + tc.get("description", "")).lower() for tc in test_cases)
    assert any("laser" in (tc.get("name", "") + " " + tc.get("description", "")).lower() for tc in test_cases)
    assert any("battery" in (tc.get("name", "") + " " + tc.get("description", "")).lower() or "lithium" in (tc.get("name", "") + " " + tc.get("description", "")).lower() for tc in test_cases)

    # Participant Tiers
    assert any("category i" in (tc.get("name", "") + " " + tc.get("description", "")).lower() or "school" in (tc.get("name", "") + " " + tc.get("description", "")).lower() for tc in test_cases)
    assert any("category ii" in (tc.get("name", "") + " " + tc.get("description", "")).lower() or "college" in (tc.get("name", "") + " " + tc.get("description", "")).lower() for tc in test_cases)
    assert any("category iii" in (tc.get("name", "") + " " + tc.get("description", "")).lower() or "industry" in (tc.get("name", "") + " " + tc.get("description", "")).lower() for tc in test_cases)


def test_adversarial_boundary_conditions(test_cases):
    """Verify specific numerical threshold pairs exist and assert boundaries."""
    boundary_cases = [tc for tc in test_cases if tc.get("case_type") == "Boundary"]
    assert len(boundary_cases) >= 8

    # Weight: 35.00kg pass vs 35.05kg DQ
    tc_022 = next((tc for tc in test_cases if tc.get("case_id") == "ROVC-TC-022"), None)
    assert tc_022 is not None
    t022_str = json.dumps(tc_022)
    assert "35.00" in t022_str and "35.05" in t022_str and ("disqualif" in t022_str.lower() or "dq" in t022_str.lower())

    # Gate: 1.00m aperture with <=5cm clearance
    tc_002 = next((tc for tc in test_cases if tc.get("case_id") == "ROVC-TC-002"), None)
    assert tc_002 is not None
    t002_str = json.dumps(tc_002)
    assert "1.00m" in t002_str and "5cm" in t002_str

    # Depth: 0.6m min and 1.8m max
    tc_003 = next((tc for tc in test_cases if tc.get("case_id") == "ROVC-TC-003"), None)
    assert tc_003 is not None
    t003_str = json.dumps(tc_003)
    assert "0.6" in t003_str and "1.8" in t003_str

    # Footprint: L<75cm, B<65.01cm, H<60.01cm
    tc_024 = next((tc for tc in test_cases if tc.get("case_id") == "ROVC-TC-024"), None)
    assert tc_024 is not None
    t024_str = json.dumps(tc_024)
    assert "75" in t024_str and "65" in t024_str and "60" in t024_str


def test_adversarial_exception_scenarios(test_cases):
    """Verify negative/failure fault handling scenarios across all dimensions."""
    exception_cases = [tc for tc in test_cases if tc.get("case_type") == "Exception"]
    assert len(exception_cases) >= 8

    # Specific critical failure modes
    ids = {tc["case_id"] for tc in exception_cases}
    assert "ROVC-TC-005" in ids  # Thruster stall
    assert "ROVC-TC-011" in ids  # Illumination failure in dark pool
    assert "ROVC-TC-012" in ids  # Wrong color balloon touch
    assert "ROVC-TC-019" in ids  # Surface broach DQ in maze
    assert "ROVC-TC-020" in ids  # Dead-end entrapment
    assert "ROVC-TC-026" in ids  # Kill switch latch failure
    assert "ROVC-TC-028" in ids  # Motor short-circuit fuse blow
    assert "ROVC-TC-032" in ids  # Oil leak abort


def test_rubric_evaluation_and_sensitivity_oracle(test_cases):
    """Directly evaluate suite with TestCaseReviewerAgent and verify sensitivity under adversarial perturbations."""
    reviewer = TestCaseReviewerAgent(llm_service=None, knowledge_service=None)
    eval_res = reviewer.evaluate_suite(test_cases)
    assert eval_res["score"] >= 85.0
    assert eval_res["score"] == 100.0
    assert eval_res["recommendation"] == "Pass"

    # Perturbation 1: Strip boundary cases -> boundary score drops
    non_bound = [tc for tc in test_cases if tc.get("case_type") != "Boundary"]
    res_bound = reviewer.evaluate_suite(non_bound)
    assert res_bound["rubric_breakdown"]["boundaries"] < 20.0
    assert res_bound["score"] < 100.0

    # Perturbation 2: Strip exception cases -> exception score drops
    non_exc = [tc for tc in test_cases if tc.get("case_type") != "Exception"]
    res_exc = reviewer.evaluate_suite(non_exc)
    assert res_exc["rubric_breakdown"]["exceptions"] < 20.0
    assert res_exc["score"] < 100.0

    # Perturbation 3: Inject duplicate IDs -> redundancy score drops
    dup_cases = copy.deepcopy(test_cases)
    dup_cases[1]["case_id"] = dup_cases[0]["case_id"]
    res_dup = reviewer.evaluate_suite(dup_cases)
    assert res_dup["rubric_breakdown"]["redundancy"] < 15.0
    assert res_dup["score"] < 100.0


def test_exporters_integrity_and_formula_injection_defense():
    """Verify Excel formula injection defense, JUnit XML parsing, and Mindmap structure."""
    # Excel inspection
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb["DTU ROVC 2026 Test Cases"]
    assert ws.max_row == 41
    assert ws.max_column == 8

    for r in range(2, 42):
        for c in range(1, 9):
            cell = ws.cell(row=r, column=c)
            assert cell.data_type == "s"
            val = str(cell.value or "")
            if val.startswith(("=", "+", "-", "@")):
                assert val.startswith("'")

    # Formula injection unit test
    assert _sanitize_excel_text("=CMD|' /C calc'!A0") == "'=CMD|' /C calc'!A0"
    assert _sanitize_excel_text("+CMD|' /C calc'!A0") == "'+CMD|' /C calc'!A0"
    assert _sanitize_excel_text("@SUM(1,2)") == "'@SUM(1,2)"
    assert _sanitize_excel_text("-5") == "'-5"
    assert _sanitize_excel_text("+44") == "'+44"
    assert _sanitize_excel_text("Safe String") == "Safe String"

    # JUnit XML inspection
    tree = ET.parse(XML_PATH)
    root = tree.getroot()
    assert root.tag == "testsuites"
    assert root.attrib["tests"] == "40"
    testcases = root.findall(".//testcase")
    assert len(testcases) == 40
    for tc in testcases:
        assert tc.attrib["classname"].startswith("CodeBrain.ROVC.")
        assert int(tc.attrib["assertions"]) >= 1
        assert tc.find("system-out") is not None

    # Mindmap inspection
    with open(MINDMAP_PATH, "r", encoding="utf-8") as f:
        mindmap_text = f.read()
    for idx in range(1, 41):
        assert f"ROVC-TC-{idx:03d}" in mindmap_text
