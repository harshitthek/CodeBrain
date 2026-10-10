import io
import json
import pytest
from django.test import RequestFactory
from django.urls import reverse

from apps.core.models import TestCase, TestCaseReview
from apps.ai_agents.test_case_generator.schemas import (
    GeneratedTestCaseItem,
    extract_and_parse_test_cases,
)
from apps.ai_agents.test_case_reviewer.views import (
    export_test_cases_excel,
    case_review,
)
from apps.core.views import update_test_case_status
from apps.ai_agents.prd_analyzer.views import extract_content_from_file
from django.core.files.uploadedfile import SimpleUploadedFile


@pytest.mark.django_db
class TestPhase1Suite:

    def test_pydantic_schema_validation(self):
        """Verify GeneratedTestCaseItem handles strings and lists gracefully."""
        # 1. Standard list inputs
        item1 = GeneratedTestCaseItem.model_validate({
            "description": "Valid login test",
            "test_steps": ["Enter username", "Enter password", "Click Submit"],
            "expected_results": ["Dashboard displayed"],
            "priority": "P1"
        })
        assert len(item1.test_steps) == 3
        assert item1.priority == "P1"

        # 2. String inputs auto-coerced to lists and list preconditions joined with newlines
        item2 = GeneratedTestCaseItem.model_validate({
            "description": "String step test",
            "preconditions": ["Precondition 1", "Precondition 2"],
            "test_steps": "Step 1\nStep 2\nStep 3",
            "expected_results": "Expected result here",
        })
        assert isinstance(item2.test_steps, list)
        assert len(item2.test_steps) == 3
        assert isinstance(item2.expected_results, list)
        assert item2.expected_results[0] == "Expected result here"
        assert item2.preconditions == "Precondition 1\nPrecondition 2"

        # 3. Empty required fields raise ValidationError
        with pytest.raises(Exception):
            GeneratedTestCaseItem.model_validate({
                "description": "   ",
                "test_steps": ["Step 1"],
                "expected_results": ["Result 1"],
            })
        with pytest.raises(Exception):
            GeneratedTestCaseItem.model_validate({
                "description": "Valid description",
                "test_steps": [],
                "expected_results": ["Result 1"],
            })


    def test_extract_and_parse_json_markdown_blocks(self):
        """Verify extraction from markdown fences, dirty preambles, and brackets."""
        markdown_response = """
Here are your requested test cases:
```json
[
  {
    "description": "Checkout test",
    "test_steps": ["Add item", "Proceed to payment"],
    "expected_results": ["Receipt generated"]
  }
]
```
Let me know if you need more!
"""
        cases = extract_and_parse_test_cases(markdown_response)
        assert len(cases) == 1
        assert cases[0]["description"] == "Checkout test"

    def test_extract_and_parse_trailing_commas(self):
        """Verify greedy bracket parser repairs trailing commas."""
        dirty_json = """
[
  {
    "description": "Dirty comma test",
    "test_steps": ["Action 1",],
    "expected_results": ["Outcome 1",],
  },
]
"""
        cases = extract_and_parse_test_cases(dirty_json)
        assert len(cases) == 1
        assert cases[0]["description"] == "Dirty comma test"

    def test_update_test_case_status_endpoint(self):
        """Verify update_test_case_status updates DB and creates review comments."""
        tc = TestCase.objects.create(
            title="Search Test",
            description="Verify search bar filters products",
            test_steps="1. Type query\n2. Press Enter",
            expected_results="Relevant products listed",
            status="pending"
        )

        rf = RequestFactory()
        payload = {
            "test_case_id": tc.id,
            "status": "approved",
            "comments": "Looks solid, approved for sprint 12."
        }
        request = rf.post(
            "/api/update-status/",
            data=json.dumps(payload),
            content_type="application/json"
        )
        response = update_test_case_status(request)
        assert response.status_code == 200
        data = json.loads(response.content)
        assert data["success"] is True

        tc.refresh_from_db()
        assert tc.status == "approved"

        # Check that review record was recorded
        reviews = TestCaseReview.objects.filter(test_case=tc)
        assert reviews.count() == 1
        assert reviews.first().review_comments == "Looks solid, approved for sprint 12."

    def test_export_test_cases_excel_openpyxl(self):
        """Verify openpyxl Excel exporter enforces auth and outputs valid .xlsx spreadsheet."""
        from django.contrib.auth.models import User

        tc = TestCase.objects.create(
            title="Export Test",
            description="Verify export to Excel",
            test_steps="Step 1",
            expected_results="Result 1",
            status="approved",
            priority="p1"
        )

        rf = RequestFactory()
        unauth_req = rf.get(f"/test_case_reviewer/api/export-test-cases-excel/?ids={tc.id}")
        unauth_resp = export_test_cases_excel(unauth_req)
        assert unauth_resp.status_code == 403

        request = rf.get(f"/test_case_reviewer/api/export-test-cases-excel/?ids={tc.id}")
        request.user = User.objects.create_superuser(username="exporter", email="exporter@example.com", password="pw")
        response = export_test_cases_excel(request)

        assert response.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response["Content-Type"]
        assert response["Content-Disposition"].endswith('.xlsx"')

        # Verify readable by openpyxl
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        assert "Test Cases" in wb.sheetnames
        sheet = wb["Test Cases"]
        assert sheet.cell(row=1, column=1).value == "No."
        assert sheet.cell(row=2, column=2).value == tc.id


    def test_prd_extract_content_from_file_text_and_md(self):
        """Verify prd analyzer file extractor handles .txt and .md files."""
        md_file = SimpleUploadedFile("spec.md", b"# Feature Specification\nUser can reset password via email.")
        content = extract_content_from_file(md_file, ".md")
        assert "# Feature Specification" in content

        txt_file = SimpleUploadedFile("notes.txt", b"Requirement notes line 1\nRequirement notes line 2")
        content_txt = extract_content_from_file(txt_file, ".txt")
        assert "Requirement notes line 1" in content_txt

    def test_prd_analyzer_multipart_file_upload(self, monkeypatch):
        """Verify prd_analyzer view processes a file-only multipart POST without RawPostDataException."""
        from apps.ai_agents.prd_analyzer.views import prd_analyzer
        from apps.ai_agents.prd_analyzer.analyser import PrdAnalyserAgent

        monkeypatch.setattr(
            PrdAnalyserAgent,
            "analyse",
            lambda self, content: {"summary": {"total_test_points": 1}, "test_points": []}
        )

        rf = RequestFactory()
        md_file = SimpleUploadedFile("prd_spec.md", b"# Checkout Flow\nUser completes payment.", content_type="text/markdown")
        request = rf.post("/prd_analyzer/", data={"single_file": md_file, "input_mode": "upload"})
        response = prd_analyzer(request)

        assert response.status_code == 200
        payload = json.loads(response.content)
        assert payload["success"] is True
        assert payload["result"]["summary"]["total_test_points"] == 1

