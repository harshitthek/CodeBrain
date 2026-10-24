"""Automated test suite for Milestone M3: Multi-Format Test Suite Exporter Integration.

Verifies that all 4 deliverables in /Users/pernavjain/Documents/Tutorial/
are strictly valid, non-zero, well-formed, and satisfy CodeBrain's enterprise contracts.
"""

import json
import xml.etree.ElementTree as ET
from pathlib import Path
import pytest
import openpyxl

from apps.ai_agents.test_case_generator.schemas import GeneratedTestCaseItem
from apps.core.exporters import (
    JUnitExporter,
    XMindExporter,
    ExcelTestSuiteExporter,
    _sanitize_excel_text,
)

TUTORIAL_DIR = (
    Path("/Users/pernavjain/Documents/Tutorial")
    if (Path("/Users/pernavjain/Documents/Tutorial") / "rovc_testcases.json").exists()
    else Path(__file__).resolve().parent.parent / "docs" / "tutorial_qa"
)
JSON_PATH = TUTORIAL_DIR / "rovc_testcases.json"
XLSX_PATH = TUTORIAL_DIR / "rovc_testcases.xlsx"
XML_PATH = TUTORIAL_DIR / "junit_testcases.xml"
MINDMAP_PATH = TUTORIAL_DIR / "testcases_mindmap.md"


class TestROVCExportsSuite:
    """Rigorous verification of Milestone M3 multi-format export deliverables."""

    def test_json_export_structure_and_telemetry(self):
        """Verify rovc_testcases.json contains metadata, quality score telemetry, and 40 valid cases."""
        assert JSON_PATH.exists(), f"JSON export not found at {JSON_PATH}"
        assert JSON_PATH.stat().st_size > 0, "JSON file is empty"

        with open(JSON_PATH, "r", encoding="utf-8") as f:
            data = json.load(f)

        assert "metadata" in data
        assert "test_cases" in data

        metadata = data["metadata"]
        assert metadata["total_test_cases"] == 40
        assert metadata["final_score"] == 100.0
        assert metadata["converged"] is True
        assert metadata["iterations_run"] == 2

        # Verify quality score telemetry
        telemetry = metadata.get("quality_score_telemetry", {})
        assert telemetry.get("final_score") == 100.0
        assert telemetry.get("target_threshold") == 85.0
        assert telemetry.get("converged") is True

        rubric = telemetry.get("iteration_2_rubric_breakdown", {})
        assert rubric.get("completeness") == 20.0
        assert rubric.get("coverage") == 25.0
        assert rubric.get("boundaries") == 20.0
        assert rubric.get("exceptions") == 20.0
        assert rubric.get("redundancy") == 15.0

        # Validate all 40 cases against GeneratedTestCaseItem schema
        cases = data["test_cases"]
        assert len(cases) == 40
        for tc in cases:
            validated = GeneratedTestCaseItem.model_validate(tc)
            assert validated.case_id.startswith("ROVC-TC-")
            assert len(validated.test_steps) >= 1
            assert len(validated.expected_results) >= 1
            assert validated.case_type in ["Functional", "Boundary", "Exception", "Security", "Performance"]
            assert validated.priority in ["High", "Medium", "Low"]

    def test_excel_export_openpyxl_and_formula_sanitization(self):
        """Verify rovc_testcases.xlsx structure, styling, cell types, and formula injection protection."""
        assert XLSX_PATH.exists(), f"Excel export not found at {XLSX_PATH}"
        assert XLSX_PATH.stat().st_size > 0, "Excel file is empty"

        wb = openpyxl.load_workbook(XLSX_PATH)
        assert "DTU ROVC 2026 Test Cases" in wb.sheetnames
        ws = wb["DTU ROVC 2026 Test Cases"]

        # 40 cases + 1 header row = 41 rows
        assert ws.max_row == 41
        assert ws.max_column == 8

        expected_headers = [
            "Case ID",
            "Name",
            "Priority",
            "Case Type",
            "Preconditions",
            "Test Steps",
            "Expected Results",
            "Category / Tag"
        ]
        actual_headers = [ws.cell(row=1, column=c).value for c in range(1, 9)]
        assert actual_headers == expected_headers

        # Verify cell data_type is 's' across all rows
        for r in range(2, 42):
            for c in range(1, 9):
                cell = ws.cell(row=r, column=c)
                assert cell.data_type == "s", f"Cell ({r},{c}) data_type is {cell.data_type}, expected 's'"
                val = str(cell.value or "")
                if val.startswith(("=", "+", "-", "@")):
                    assert val.startswith("'"), f"Cell ({r},{c}) failed formula sanitization: {val}"

        # Unit test _sanitize_excel_text helper directly
        assert _sanitize_excel_text("=SUM(1,2)") == "'=SUM(1,2)"
        assert _sanitize_excel_text("+CMD|' /C calc'!A0") == "'+CMD|' /C calc'!A0"
        assert _sanitize_excel_text("-10+5") == "'-10+5"
        assert _sanitize_excel_text("@HYPERLINK('http://evil.com')") == "'@HYPERLINK('http://evil.com')"
        assert _sanitize_excel_text("   =DANGEROUS") == "'   =DANGEROUS"
        assert _sanitize_excel_text("Safe Text") == "Safe Text"

    def test_junit_xml_export_parsing_and_schema(self):
        """Verify junit_testcases.xml well-formed XML, 40 testcases, classnames, and assertions."""
        assert XML_PATH.exists(), f"JUnit XML not found at {XML_PATH}"
        assert XML_PATH.stat().st_size > 0, "JUnit XML file is empty"

        tree = ET.parse(XML_PATH)
        root = tree.getroot()

        assert root.tag == "testsuites"
        assert root.attrib.get("tests") == "40"
        assert root.attrib.get("failures") == "0"
        assert root.attrib.get("errors") == "0"

        testcases = root.findall(".//testcase")
        assert len(testcases) == 40

        for tc in testcases:
            assert "classname" in tc.attrib
            assert tc.attrib["classname"].startswith("CodeBrain.ROVC.")
            assert "name" in tc.attrib
            assert "id" in tc.attrib
            assert tc.attrib["id"].startswith("ROVC-TC-")
            assert "assertions" in tc.attrib
            assert int(tc.attrib["assertions"]) >= 1

            # Verify system-out with assertions
            sys_out = tc.find("system-out")
            assert sys_out is not None
            assert "Verifiable Assertions:" in (sys_out.text or "")

    def test_markdown_mindmap_hierarchy_and_coverage(self):
        """Verify testcases_mindmap.md Markmap hierarchy covers Domain -> Task 1-3/Safety/Stories -> Subsystems."""
        assert MINDMAP_PATH.exists(), f"MindMap not found at {MINDMAP_PATH}"
        assert MINDMAP_PATH.stat().st_size > 0, "MindMap file is empty"

        with open(MINDMAP_PATH, "r", encoding="utf-8") as f:
            content = f.read()

        assert "# DTU ROVC 2026 Competition Test Suite" in content
        assert "## Task 1: Qualification Round Submerged Gate Navigation" in content
        assert "## Task 2: Low-Visibility Vision Target Identification and Touch" in content
        assert "## Task 3: The Maze Walk Submerged Traversal" in content
        assert "## Physical & Safety Constraints" in content
        assert "## User Stories & Governance" in content

        # Verify all 40 test case IDs are present in the mindmap
        for idx in range(1, 41):
            cid = f"ROVC-TC-{idx:03d}"
            assert cid in content, f"Missing test case {cid} in mindmap"

        # Verify key subsections exist
        assert "### Subsystem: Gate Aperture & Submerged Navigation" in content
        assert "### Subsystem: Low-Visibility Vision & Onboard Illumination" in content
        assert "### Subsystem: Maze Channel Navigation & Obstacle Avoidance" in content
        assert "### Subsystem: Vehicle Mass & Compact Footprint Envelope" in content
        assert "### Subsystem: External Kill Switch & Emergency Cutoff" in content
        assert "### Subsystem: Category I — School Tier Teleoperation & Safety" in content
